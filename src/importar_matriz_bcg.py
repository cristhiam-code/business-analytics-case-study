# =====================================================
# Script: importar_matriz_bcg.py
# Objetivo: leer MatrizBCG2_limpio_v2.xlsx y cargarlo
# como tabla "matriz_bcg" dentro de farmanorte.db
# =====================================================

import openpyxl
import sqlite3

from utilidades import CARPETA_RAW, CARPETA_DATABASE, ARCHIVO_DB, asegurar_directorios

# --- Rutas del proyecto, no rutas absolutas del equipo ---
XLSX_PATH = CARPETA_RAW / "MatrizBCG2_limpio_v2.xlsx"
DB_PATH = ARCHIVO_DB


def crear_tabla(conn):
    """
    Borra la tabla matriz_bcg si ya existia y la vuelve a crear vacia.
    """
    conn.execute('DROP TABLE IF EXISTS "matriz_bcg"')
    conn.execute('''
        CREATE TABLE "matriz_bcg" (
            "codigo_barras" TEXT PRIMARY KEY,
            "producto" TEXT,
            "laboratorio" TEXT,
            "existencia" INTEGER,
            "precio_venta" REAL,
            "costo_promedio" REAL,
            "unid_vendidas" INTEGER,
            "ranking_1" INTEGER,
            "val_vendido" REAL,
            "ranking_2" INTEGER,
            "rentabilidad" REAL,
            "ranking_3" INTEGER,
            "suma_ranking" INTEGER,
            "ranking_final" INTEGER,
            "grupo" TEXT
        )
    ''')


def importar(conn):
    """
    Abre el Excel, lee todas las filas de datos (saltando el encabezado)
    y las inserta de una sola vez en matriz_bcg.
    """
    if not XLSX_PATH.exists():
        raise FileNotFoundError(f"No se encontró el archivo Excel: {XLSX_PATH}")

    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True)
    hoja = wb[wb.sheetnames[0]]

    filas = []

    for fila in hoja.iter_rows(min_row=2, values_only=True):
        filas.append(fila)

    conn.executemany(
        'INSERT INTO "matriz_bcg" VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
        filas
    )

    print(f"matriz_bcg: {len(filas)} filas importadas")


def main():
    """
    Orquesta todo el proceso en orden: crear carpeta, conectar, crear tabla,
    importar datos, guardar cambios y cerrar.
    """
    asegurar_directorios(CARPETA_DATABASE)
    conn = sqlite3.connect(DB_PATH)

    crear_tabla(conn)
    importar(conn)

    conn.commit()
    conn.close()

    print("Listo.")


if __name__ == "__main__":
    main()